"""Lógica do gerador de vale presente (cruzamento apuração x base de cadastro).

Baseado no gerar_todos_vales.py do Núcleo. Diferenças: funciona com caminhos
informados (sem pastas fixas), lê só a aba de premiação em modo streaming e
devolve um resumo estruturado para a interface web.
"""
import os
import unicodedata

import pandas as pd
from openpyxl import load_workbook

ABA_BASE = "BASE COM CADASTRO"
LINHA_CAB = 4
MAX_COLS = 12

CATEGORIAS = {
    "OPERADOR", "SUPERVISOR", "GESTOR", "COORDENADOR", "COORDENADOR ALGAR",
    "GERENTE", "AGENTE", "CANAL DEALER", "CANAL DIGITAL", "COORDENADOR CALLINK",
    "GESTOR CALLINK", "SUPERVISOR CALLINK", "GERENTE CALLINK", "CANAL CALLINK",
    "CANAL REMOTO",
}
CARGOS = ("COORDENADOR", "GESTOR", "SUPERVISOR", "GERENTE", "CANAL")
IGNORAR_NOME = {"", "GANHADORES", "NOME", "PREMIACOES", "PREMIACAO", "CANAL", "TIPO"}

# (coluna do nome, coluna do valor), índices base 0, sem repetição
PARES_COLUNAS = [(1, 4), (6, 9), (1, 5), (7, 11), (2, 5), (8, 11), (7, 10)]


def norm(s):
    if s is None or (not isinstance(s, str) and pd.isna(s)):
        return ""
    s = " ".join(str(s).strip().upper().split())
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def fmt_cpf(x):
    if x is None or (not isinstance(x, str) and pd.isna(x)):
        return ""
    x = str(x).strip()
    if x.endswith(".0"):
        x = x[:-2]
    x = x.replace(".", "").replace("-", "").replace(" ", "")
    if x.isdigit() and 1 <= len(x) <= 11:
        x = x.zfill(11)
    return x


def nome_saida(nome_arquivo, pasta_saida):
    base = os.path.splitext(os.path.basename(nome_arquivo))[0]
    base = base.replace("Apuração ", "").replace("Apuracao ", "")
    base = " ".join(base.split())
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "-")
    return os.path.join(pasta_saida, f"ValePresente {base}.xlsx")


def ler_base(caminho):
    """Lê a BASE COMPLETA e devolve {nome_normalizado: registro}."""
    try:
        df = pd.read_excel(caminho, sheet_name=ABA_BASE, header=LINHA_CAB - 1, dtype=str)
    except ValueError:
        raise ValueError(f"A base precisa ter uma aba chamada '{ABA_BASE}'.")
    ordem = ["NOME", "CPF", "DATA_NASC", "NOME_MAE", "EMAIL", "EMBOSSING", "VALOR"]
    df = df.rename(columns={c: ordem[i] for i, c in enumerate(df.columns) if i < len(ordem)})
    if "NOME" not in df.columns:
        raise ValueError("Base sem coluna de nome (esperado: *NOME COMPLETO na linha 4).")
    df["NOME"] = df["NOME"].astype(str)
    df = df[df["NOME"].str.strip() != ""]
    df = df[~df["NOME"].str.upper().isin(["*NOME COMPLETO", "NAN"])]
    df = df[df["NOME"].str.len() > 2].copy()
    df["NORM"] = df["NOME"].apply(norm)
    df["CPF"] = df["CPF"].apply(fmt_cpf)
    df = df.drop_duplicates(subset=["NORM"], keep="first")
    return {r["NORM"]: r for _, r in df.iterrows()}


def escolher_aba(caminho):
    wb = load_workbook(caminho, read_only=True, data_only=True)
    abas = wb.sheetnames
    wb.close()
    for a in abas:
        na = norm(a)
        if "PREMIAC" in na or "PREMIAD" in na:
            return a, abas
    return None, abas


def extrair_pares(caminho, aba):
    """Lista de (nome, valor_inteiro) das linhas da aba de premiação."""
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[aba]
    pares = []
    for row in ws.iter_rows(min_col=1, max_col=MAX_COLS, values_only=True):
        vals = ["" if v is None else str(v).strip() for v in row]
        vals += [""] * (MAX_COLS - len(vals))
        for cn, cv in PARES_COLUNAS:
            nome = vals[cn]
            valor = vals[cv].replace(" ", "").replace(",", ".")
            nu = norm(nome)
            if nu in IGNORAR_NOME or "SEM PREMIADO" in nu or "TOTAL" in nu:
                continue
            if valor == "":
                continue
            try:
                v = float(valor.replace("R$", ""))
            except ValueError:
                continue
            if v <= 0:
                continue
            if "LTDA" in nu or "EIRELI" in nu or "S.A" in nu:
                continue
            if len([p for p in nu.split() if len(p) > 1]) < 2:
                continue
            if nu in CATEGORIAS:
                continue
            if "CALLINK" in nu and any(x in nu for x in CARGOS):
                continue
            pares.append((nome, int(v)))
    wb.close()
    return pares


def _linha(b, valor):
    return [b["NOME"], b["CPF"], b.get("DATA_NASC", ""), b.get("NOME_MAE", ""),
            b.get("EMAIL", ""), b.get("EMBOSSING", ""), valor]


def processar_apuracao(caminho_apuracao, nome_original, dic, caminho_base, pasta_saida):
    """Gera o ValePresente de uma apuração. Devolve dict com o resumo."""
    aba, abas = escolher_aba(caminho_apuracao)
    if aba is None:
        return {"arquivo": nome_original, "ok": False,
                "erro": f"Não achei aba de premiação (abas: {', '.join(abas)})"}

    pares = extrair_pares(caminho_apuracao, aba)
    grupo = {}
    for nome, v in pares:
        k = norm(nome)
        grupo.setdefault(k, {"nome": nome, "valor": 0})["valor"] += v

    com, sem, aproximados = [], [], []
    for k, g in sorted(grupo.items()):
        if k in dic:
            com.append(_linha(dic[k], g["valor"]))
            continue
        achou = None
        if k and len(k) > 8:
            for kb, b in dic.items():
                if k in kb or kb in k:
                    achou = b
                    break
        if achou is not None:
            com.append(_linha(achou, g["valor"]))
            aproximados.append({"apuracao": g["nome"], "base": achou["NOME"], "valor": g["valor"]})
        else:
            sem.append([g["nome"], "", "", "", "", "", g["valor"]])

    saida = nome_saida(nome_original, pasta_saida)
    wb = load_workbook(caminho_base)
    for an, linhas in (("COM CADASTRO", com), ("SEM CADASTRO", sem)):
        if an not in wb.sheetnames:
            raise ValueError(f"A base precisa ter a aba '{an}' (modelo do cartão).")
        ws = wb[an]
        for r in range(5, ws.max_row + 1):
            for c in range(1, 8):
                ws.cell(row=r, column=c).value = None
        for i, lin in enumerate(linhas):
            for c, val in enumerate(lin, start=1):
                cell = ws.cell(row=5 + i, column=c)
                if isinstance(val, str) and val.strip().lower() in ("nan", ""):
                    val = None
                if c == 7 and val is not None:
                    cell.value = int(val)
                    cell.number_format = "0"
                else:
                    cell.value = val
    wb.save(saida)

    return {
        "arquivo": nome_original, "ok": True, "aba": aba, "saida": saida,
        "linhas": len(pares), "pessoas": len(grupo),
        "com": len(com), "sem": len(sem),
        "total": sum(g["valor"] for g in grupo.values()),
        "total_com": sum(l[6] for l in com), "total_sem": sum(l[6] for l in sem),
        "sem_lista": [{"nome": l[0], "valor": l[6]} for l in sem],
        "aproximados": aproximados,
    }


def consolidar_sem_cadastro(resultados, caminho):
    """Planilha única com todos os SEM CADASTRO de todas as apurações."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "SEM CADASTRO"
    ws.append(["Apuração", "Nome", "Valor"])
    n = 0
    for r in resultados:
        for p in r.get("sem_lista", []):
            ws.append([r["arquivo"], p["nome"], p["valor"]])
            n += 1
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 45
    wb.save(caminho)
    return n
